import openpyxl
import PyPDF2
from docx import Document
from bs4 import BeautifulSoup
from curl_cffi import requests

def pdf_extract_text(file_path):
    full_text = ""
    with open(file_path, "rb") as file:
        reader = PyPDF2.PdfReader(file)
        for page in reader.pages:
            full_text += page.extract_text() + "\n"
    return full_text

def docx_extract_text(file_path):
    doc = Document(file_path)
    full_text = []
    for paragraph in doc.paragraphs:
        full_text.append(paragraph.text)
    return "\n".join(full_text)

def excel_extract_text(file_path):
    workbook = openpyxl.load_workbook(file_path)
    
    sheet = workbook.active
    
    text_lines = []
    
    for row in sheet.iter_rows(values_only=True):
        line = '\t'.join([str(cell) if cell is not None else '' for cell in row])
        text_lines.append(line)
    
    result = '\n'.join(text_lines)
    return result

def web_extract_text(resource):
    response = requests.get(resource)
    html_content = response.text

    soup = BeautifulSoup(html_content, 'html.parser')
    inner_text = soup.get_text()
    cleaned_text = ' '.join(inner_text.split())
    return cleaned_text